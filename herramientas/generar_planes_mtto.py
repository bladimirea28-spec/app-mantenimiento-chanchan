"""Genera datos/planes-mantenimiento.js y plantillas/CONTROL_MANTENIMIENTO.xlsx
a partir del Excel de control de mantenimiento.

Uso (desde la carpeta del proyecto):
    python herramientas/generar_planes_mtto.py "C:/Users/Bladimir/Desktop/CONTROL DE MANTENIMIENTO - CHAN CHANNN.xlsx"

Vuelve a ejecutarlo cada vez que cambies equipos, planes, materiales,
inspecciones o ciclos en el Excel, y publica la app.
"""
import json
import shutil
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent


def texto(valor):
    return "" if valor is None else str(valor).strip()


def numero(valor):
    try:
        return float(valor)
    except (TypeError, ValueError):
        return 0.0


def main(ruta_excel):
    libro = openpyxl.load_workbook(ruta_excel, data_only=True)

    equipos = []
    for fila in libro["EQUIPOS"].iter_rows(min_row=3, values_only=True):
        if not fila[1]:
            continue
        equipos.append({
            "codigo": texto(fila[1]), "descripcion": texto(fila[2]), "area": texto(fila[3]),
            "criticidad": texto(fila[4]), "plan": texto(fila[5]), "control": texto(fila[6]),
            "intervalo": numero(fila[7]),
        })

    planes = OrderedDict()

    def plan(nombre):
        return planes.setdefault(nombre, {"ciclo": [], "materiales": OrderedDict(), "inspeccion": OrderedDict()})

    for fila in libro["BD_CICLO"].iter_rows(min_row=3, values_only=True):
        if not fila[1]:
            continue
        plan(texto(fila[1]))["ciclo"].append({
            "orden": int(numero(fila[2])), "paso": texto(fila[3]),
            "lectura": numero(fila[4]), "mantenimiento": texto(fila[5]),
        })

    for fila in libro["BD_PLAN"].iter_rows(min_row=3, values_only=True):
        if not fila[1]:
            continue
        plan(texto(fila[1]))["materiales"].setdefault(texto(fila[2]), []).append({
            "sistema": texto(fila[3]), "descripcion": texto(fila[4]), "cantidad": numero(fila[5]),
            "unidad": texto(fila[6]), "precio": numero(fila[7]),
        })

    for fila in libro["BD_INSPECCION"].iter_rows(min_row=3, values_only=True):
        if not fila[1]:
            continue
        plan(texto(fila[1]))["inspeccion"].setdefault(texto(fila[2]), []).append({
            "sistema": texto(fila[3]), "actividad": texto(fila[4]), "tarea": texto(fila[5]),
        })

    for p in planes.values():
        p["ciclo"].sort(key=lambda x: x["orden"])

    destino = RAIZ / "datos" / "planes-mantenimiento.js"
    with open(destino, "w", encoding="utf-8") as f:
        f.write("/* ============ PLANES DE MANTENIMIENTO PREVENTIVO ============\n")
        f.write("   ARCHIVO GENERADO: no lo edites a mano. Sale del Excel\n")
        f.write("   \"CONTROL DE MANTENIMIENTO\" con herramientas/generar_planes_mtto.py\n")
        f.write("   (hojas EQUIPOS, BD_CICLO, BD_PLAN y BD_INSPECCION). */\n\n")
        f.write("const EQUIPOS_MTTO = " + json.dumps(equipos, ensure_ascii=False, indent=1) + ";\n\n")
        f.write("const PLANES_MTTO = " + json.dumps(planes, ensure_ascii=False) + ";\n")

    plantilla = RAIZ / "plantillas" / "CONTROL_MANTENIMIENTO.xlsx"
    plantilla.parent.mkdir(exist_ok=True)
    shutil.copyfile(ruta_excel, plantilla)

    print(f"Equipos: {len(equipos)} | Planes: {len(planes)}")
    for nombre, p in planes.items():
        print(f"  {nombre}: {len(p['ciclo'])} pasos, {sum(len(v) for v in p['materiales'].values())} materiales, "
              f"{sum(len(v) for v in p['inspeccion'].values())} tareas")
    print("Generado:", destino)
    print("Plantilla:", plantilla)


if __name__ == "__main__":
    main(sys.argv[1])
