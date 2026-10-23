"""Define las tablas que se copian a Dataverse (modo espejo) y genera:
  1) Un Excel por tabla en el Escritorio (carpeta "DATAVERSE CHAN CHAN"), con
     los encabezados y una fila de ejemplo ficticia, para crear cada tabla en
     Power Apps con "Crear tabla desde Excel".
  2) apps-script/espejo_config.gs.txt: la configuración que usa Codigo.gs.

Uso: python herramientas/tablas_dataverse.py
"""
import json
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

RAIZ = Path(__file__).resolve().parent.parent
SALIDA_EXCEL = Path.home() / "Desktop" / "DATAVERSE CHAN CHAN"

FECHA = "fecha"
NUM = "numero"
TXT = "texto"
EJEMPLO = {FECHA: datetime(2026, 10, 18, 8, 30), NUM: 1250.5, TXT: "Ejemplo"}

# hoja de Google Sheets, nombre de la tabla en Dataverse, columnas que forman
# la clave única, y columnas (nombre en la hoja, tipo). La primera columna de
# cada tabla en Dataverse es CLAVE_APP (la clave única, texto).
TABLAS = [
    ("ACTIVIDADES_MTTO", "Mtto Actividades", ["ID_UNICO|ID_OT,DNI,FECHA_SYNC"], [
        ("ID_OT", TXT), ("FECHA", FECHA), ("AÑO", NUM), ("USUARIO", TXT), ("DNI", TXT), ("EQUIPO", TXT), ("AREA", TXT),
        ("CRITICIDAD", TXT), ("TIPO_MTTO", TXT), ("ESPECIALIDAD", TXT), ("DESCRIPCION", TXT), ("FOTOS", TXT),
        ("FECHA_SYNC", FECHA), ("TECNICO_RESPONSABLE", TXT), ("FECHA_HORA_INICIO", TXT), ("FECHA_HORA_FIN", TXT),
        ("TOTAL_HORAS", NUM), ("HOROMETRO_KM", TXT), ("DETALLE_SOLICITUD", TXT), ("INTEGRANTES", TXT), ("COMENTARIOS", TXT)]),
    ("HISTORIAL_FALLAS", "Mtto Fallas", ["ID_UNICO|ID_FALLA,DNI,FECHA_SYNC"], [
        ("ID_FALLA", TXT), ("FECHA", FECHA), ("TECNICO", TXT), ("DNI", TXT), ("EQUIPO", TXT), ("AREA", TXT), ("TIPO_FALLA", TXT),
        ("DESCRIPCION", TXT), ("CAUSA_RAIZ", TXT), ("ACCION_CORRECTIVA", TXT), ("HORA_INICIO", TXT), ("HORA_FIN", TXT),
        ("MINUTOS_PARADA", NUM), ("N_OPERARIOS", TXT), ("TURNO", TXT), ("TIPO_ATENCION", TXT), ("FOTOS", TXT),
        ("FECHA_SYNC", FECHA), ("TECNICO_RESPONSABLE", TXT), ("CRITICIDAD", TXT), ("ESTADO", TXT), ("RESUELTO_POR", TXT),
        ("FECHA_RESOLUCION", FECHA)]),
    ("REQUERIMIENTOS", "Mtto Requerimientos", ["ID"], [
        ("ID", TXT), ("CORRELATIVO", TXT), ("FECHA", FECHA), ("TECNICO", TXT), ("DNI", TXT), ("MOTIVO", TXT),
        ("MATERIALES", TXT), ("TOTAL", NUM), ("COMENTARIOS", TXT), ("FECHA_SYNC", FECHA), ("PRIORIDAD", TXT)]),
    ("REQ_DETALLE", "Mtto Requerimiento Detalle", ["CORRELATIVO,#"], [
        ("CORRELATIVO", TXT), ("FECHA", FECHA), ("TECNICO", TXT), ("MOTIVO", TXT), ("CANTIDAD", NUM), ("UM", TXT),
        ("CODIGO_PRODUCTO", TXT), ("NOMBRE", TXT), ("ESPECIFICACIONES", TXT), ("CODIGO_CECO", TXT), ("MAQUINARIA", TXT),
        ("PRECIO_UNITARIO", NUM), ("PRECIO_TOTAL", NUM)]),
    ("SOLICITUDES_ACTIVIDAD", "Mtto Solicitudes", ["ID_SOLICITUD"], [
        ("ID_SOLICITUD", TXT), ("FECHA", FECHA), ("SOLICITADO_POR", TXT), ("DNI", TXT), ("AREA", TXT), ("EQUIPO", TXT),
        ("PRIORIDAD", TXT), ("FECHA_TENTATIVA", TXT), ("DETALLE", TXT), ("COMENTARIOS", TXT), ("FOTOS", TXT),
        ("ESTADO", TXT), ("FECHA_SYNC", FECHA), ("ORIGEN", TXT)]),
    ("CHECKLISTS", "Mtto Checklists", ["ID_CHECKLIST"], [
        ("ID_CHECKLIST", TXT), ("FECHA", FECHA), ("HORA", TXT), ("SEMANA", TXT), ("DIA", TXT), ("PLANTILLA", TXT),
        ("CODIGO_FORMATO", TXT), ("EQUIPO_CODIGO", TXT), ("EQUIPO", TXT), ("EQUIPO_CLAVE", TXT), ("AREA", TXT),
        ("TURNO", TXT), ("TRACTOR", TXT), ("LECTURA", NUM), ("LECTURA_TIPO", TXT), ("NC", NUM), ("OBSERVADOS", NUM),
        ("PUEDE_UTILIZAR", TXT), ("OBSERVACIONES", TXT), ("FOTOS", TXT), ("USUARIO", TXT), ("DNI", TXT), ("FECHA_SYNC", FECHA)]),
    ("CHECKLIST_DETALLE", "Mtto Checklist Detalle", ["ID_CHECKLIST,ITEM"], [
        ("ID_CHECKLIST", TXT), ("FECHA", FECHA), ("SEMANA", TXT), ("PLANTILLA", TXT), ("EQUIPO_CLAVE", TXT), ("EQUIPO", TXT),
        ("SECCION", TXT), ("ITEM", TXT), ("MARCA", TXT), ("OBSERVACION", TXT)]),
    ("CHECKLIST_VISTO_BUENO", "Mtto Checklist Visto Bueno", ["PLANTILLA,SEMANA,EQUIPO_CLAVE,TIPO"], [
        ("PLANTILLA", TXT), ("SEMANA", TXT), ("EQUIPO_CLAVE", TXT), ("EQUIPO", TXT), ("TIPO", TXT), ("NOMBRE", TXT),
        ("DNI", TXT), ("FECHA", FECHA)]),
    ("MANTENIMIENTOS", "Mtto Control Mantenimientos", ["ID_MANTENIMIENTO"], [
        ("ID_MANTENIMIENTO", TXT), ("FECHA", FECHA), ("CODIGO", TXT), ("EQUIPO", TXT), ("PLAN", TXT), ("CONTROL", TXT),
        ("LECTURA", NUM), ("TIPO", TXT), ("MANTENIMIENTO_REALIZADO", TXT), ("N_OT", TXT), ("TECNICO", TXT), ("DNI", TXT),
        ("HORA_INICIO", TXT), ("HORA_FIN", TXT), ("COSTO", NUM), ("CAMBIOS_FUERA_PLAN", TXT), ("OBSERVACIONES", TXT),
        ("TRABAJO_REALIZADO", TXT), ("FALLA_RELACIONADA", TXT), ("NO_OK", NUM), ("FOTOS", TXT), ("VB_NOMBRE", TXT),
        ("VB_FECHA", FECHA), ("FECHA_SYNC", FECHA), ("REPUESTOS_TEXTO", TXT), ("REQUERIMIENTO", TXT)]),
    ("MTTO_MATERIALES", "Mtto Control Materiales", ["ID_MANTENIMIENTO,#"], [
        ("ID_MANTENIMIENTO", TXT), ("N_OT", TXT), ("FECHA", FECHA), ("CODIGO", TXT), ("TIPO", TXT), ("MANTENIMIENTO", TXT),
        ("SISTEMA", TXT), ("DESCRIPCION", TXT), ("CANTIDAD_PLAN", NUM), ("UNIDAD", TXT), ("CAMBIADO", TXT),
        ("CANTIDAD_USADA", NUM), ("PRECIO_UNIT", NUM), ("PRECIO_TOTAL", NUM)]),
    ("MTTO_INSPECCION", "Mtto Control Inspeccion", ["ID_MANTENIMIENTO,#"], [
        ("ID_MANTENIMIENTO", TXT), ("N_OT", TXT), ("FECHA", FECHA), ("CODIGO", TXT), ("MANTENIMIENTO", TXT),
        ("SISTEMA", TXT), ("ACTIVIDAD", TXT), ("TAREA", TXT), ("RESULTADO", TXT), ("OBSERVACION", TXT)]),
    ("MTTO_PARAMETROS", "Mtto Control Punto Partida", ["CODIGO"], [
        ("CODIGO", TXT), ("EQUIPO", TXT), ("LECTURA_INICIAL", NUM), ("FECHA_INICIAL", FECHA), ("ULTIMO_PM", TXT),
        ("LECTURA_ULTIMO_PM", NUM), ("FECHA_ULTIMO_PM", FECHA), ("AVISO", NUM), ("USUARIO", TXT), ("FECHA_ACTUALIZACION", FECHA)]),
    ("USUARIOS", "Mtto Usuarios", ["SLOT_ID"], [
        ("SLOT_ID", TXT), ("ROL", TXT), ("NOMBRE", TXT), ("CARGO", TXT), ("DNI", TXT), ("FECHA_REGISTRO", FECHA),
        ("AREA", TXT), ("WHATSAPP", TXT), ("CORREO", TXT), ("ES_JEFE_AREA", TXT)]),
]


def nombre_columna_dataverse(nombre_hoja):
    """Dataverse no se lleva bien con Ñ ni tildes en los nombres internos."""
    return nombre_hoja.replace("Ñ", "NI")


# Entorno personal de Bladimir (org9a7207d6): prefijo del editor y nombre del
# "conjunto" de cada tabla en la API (Dataverse lo pluraliza a su manera).
PREFIJO = "cra58_"
CONJUNTOS = {
    "Mtto Actividades": "cra58_mtto_actividadeses",
    "Mtto Fallas": "cra58_mtto_fallases",
    "Mtto Requerimientos": "cra58_mtto_requerimientoses",
    "Mtto Requerimiento Detalle": "cra58_mtto_requerimiento_detalles",
    "Mtto Solicitudes": "cra58_mtto_solicitudeses",
    "Mtto Checklists": "cra58_mtto_checklistses",
    "Mtto Checklist Detalle": "cra58_mtto_checklist_detalles",
    "Mtto Checklist Visto Bueno": "cra58_mtto_checklist_visto_buenos",
    "Mtto Control Mantenimientos": "cra58_mtto_control_mantenimientoses",
    "Mtto Control Materiales": "cra58_mtto_control_materialeses",
    "Mtto Control Inspeccion": "cra58_mtto_control_inspeccions",
    "Mtto Control Punto Partida": "cra58_mtto_control_punto_partidas",
    "Mtto Usuarios": "cra58_mtto_usuarioses",
}


def main():
    SALIDA_EXCEL.mkdir(parents=True, exist_ok=True)
    config = []
    for hoja, tabla, claves, columnas in TABLAS:
        libro = openpyxl.Workbook()
        ws = libro.active
        ws.title = tabla[:31]
        encabezados = ["CLAVE_APP"] + [nombre_columna_dataverse(c) for c, _ in columnas]
        ws.append(encabezados)
        ws.append(["EJEMPLO-001"] + [EJEMPLO[t] for _, t in columnas])
        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F4D2C")
        for i, (_, t) in enumerate(columnas, start=2):
            if t == FECHA:
                ws.cell(row=2, column=i).number_format = "dd/mm/yyyy hh:mm"
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18
        ws.column_dimensions["A"].width = 24
        archivo = SALIDA_EXCEL / f"{tabla}.xlsx"
        libro.save(archivo)
        config.append({
            "hoja": hoja, "tabla": CONJUNTOS[tabla], "clave": claves[0],
            "columnas": [[c, PREFIJO + nombre_columna_dataverse(c).lower(), t] for c, t in columnas],
        })
        print(f"{tabla:32} {len(columnas) + 1:>3} columnas -> {archivo.name}")

    destino = RAIZ / "apps-script" / "espejo_config.gs.txt"
    destino.write_text("var ESPEJO_TABLAS = " + json.dumps(config, ensure_ascii=False, indent=1) + ";\n", encoding="utf-8")
    print("Configuración:", destino)


if __name__ == "__main__":
    main()
